"""
Resume text extraction.

PDF  -> pdfplumber (primary, keeps layout better) with PyPDF2 as a fallback
DOCX -> python-docx (paragraphs + tables)
"""
import os

import pdfplumber
from PyPDF2 import PdfReader
from docx import Document


def extract_text_from_pdf(path: str) -> str:
    text_parts: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
    except Exception:
        text_parts = []

    text = "\n".join(text_parts).strip()
    if text:
        return text

    # Fallback: PyPDF2 (works on some PDFs where pdfplumber returns nothing)
    try:
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception:
        return ""


def extract_text_from_docx(path: str) -> str:
    document = Document(path)
    parts = [para.text for para in document.paragraphs]
    for table in document.tables:                 # skills are often inside tables
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(p for p in parts if p.strip()).strip()


def extract_text(path: str) -> str:
    """Dispatch on file extension. Raises ValueError for unsupported types."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext == ".docx":
        return extract_text_from_docx(path)
    raise ValueError(f"Unsupported resume format: {ext}")
